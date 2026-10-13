"""Build an editable PowerPoint counterpart for an authored 16:9 SVG figure.

Used for storing PPTX binary sources in GitHub. For custom future revisions,
edit both source scene and SVG; HTML consumes the SVG while PowerPoint remains
editable (individual rectangles, texts and connectors).
"""
from __future__ import annotations

import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

SVG_SCALE = 90.0
FONT = "Noto Sans CJK KR"

def color(raw: str) -> RGBColor:
    s = (raw or "#000000").lstrip("#")
    if len(s) != 6:
        return RGBColor(0, 0, 0)
    return RGBColor(int(s[:2], 16), int(s[2:4], 16), int(s[4:], 16))

def n(v: str, default=0.) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return default

def make(svg_path: Path, pptx_path: Path):
    root = ET.parse(svg_path).getroot()
    doc = Presentation()
    doc.slide_width = Inches(1200 / SVG_SCALE)
    doc.slide_height = Inches(675 / SVG_SCALE)
    slide = doc.slides.add_slide(doc.slide_layouts[6])
    for item in root:
        tag = item.tag.split("}")[-1]
        if tag == "rect":
            x,y,w,h = (n(item.get(a))/SVG_SCALE for a in ("x","y","width","height"))
            if not (0 <= x < 13.34 and 0 <= y < 7.51 and w > .001 and h > .001):
                continue
            rounded = n(item.get("rx")) > 0
            sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                        Inches(x), Inches(y), Inches(w), Inches(h))
            sh.fill.solid()
            sh.fill.fore_color.rgb = color(item.get("fill"))
            if item.get("stroke","none") == "none":
                sh.line.fill.background()
            else:
                sh.line.color.rgb = color(item.get("stroke"))
                sh.line.width = Pt(max(.05,n(item.get("stroke-width"),.7)*72/SVG_SCALE))
            if rounded:
                try: sh.adjustments[0] = 0.08
                except (IndexError, TypeError): pass
        elif tag == "line":
            x1,y1,x2,y2 = (n(item.get(a))/SVG_SCALE for a in ("x1","y1","x2","y2"))
            sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
            sh.line.color.rgb = color(item.get("stroke","#2F7E97"))
            sh.line.width = Pt(n(item.get("stroke-width"),2)*72/SVG_SCALE)
        elif tag == "polygon":
            # All authored polygon items are arrowheads pointing right.
            pairs = [tuple(map(float,p.split(","))) for p in item.get("points","").split() if "," in p]
            if not pairs: continue
            minx,maxx=min(p[0] for p in pairs),max(p[0] for p in pairs)
            miny,maxy=min(p[1] for p in pairs),max(p[1] for p in pairs)
            sh=slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                      Inches(minx/SVG_SCALE), Inches(miny/SVG_SCALE),
                                      Inches((maxx-minx)/SVG_SCALE), Inches((maxy-miny)/SVG_SCALE))
            sh.rotation=90
            sh.fill.solid();sh.fill.fore_color.rgb=color(item.get("fill"))
            sh.line.fill.background()
        elif tag == "text":
            value="".join(item.itertext())
            if not value.strip(): continue
            fontsize_px=n(item.get("font-size"),19)
            fontsize_pt=fontsize_px*72/SVG_SCALE
            ax=n(item.get("x"))/SVG_SCALE
            baseline=n(item.get("y"))/SVG_SCALE
            # Per-character estimated span creates editable text with adequate room.
            weighted=sum(1.0 if ord(ch)>127 else .59 for ch in value)
            width=min(12.7,max(.45,weighted*fontsize_pt/72*1.05))
            anchor=item.get("text-anchor","middle")
            x=ax if anchor=="start" else (ax-width if anchor=="end" else ax-width/2)
            x=max(0,min(13.333-width,x))
            h=max(.28,fontsize_pt/72*1.65)
            y=max(0,baseline-fontsize_pt/72*1.16)
            t=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(width),Inches(h))
            tf=t.text_frame
            tf.clear();tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            para=tf.paragraphs[0];para.text=value
            para.alignment={"start":PP_ALIGN.LEFT,"middle":PP_ALIGN.CENTER,"end":PP_ALIGN.RIGHT}.get(anchor,PP_ALIGN.CENTER)
            run=para.runs[0];run.font.name=FONT;run.font.size=Pt(fontsize_pt)
            run.font.bold=int(item.get("font-weight","400"))>=650
            run.font.color.rgb=color(item.get("fill","#213A4C"))
    pptx_path.parent.mkdir(parents=True,exist_ok=True)
    doc.save(pptx_path)
    print("EDITABLE_PPTX",pptx_path,"shapes",len(slide.shapes))

if __name__ == "__main__":
    if len(sys.argv)!=3:
        raise SystemExit("usage: svg_to_editable_pptx.py INPUT.svg OUTPUT.pptx")
    make(Path(sys.argv[1]),Path(sys.argv[2]))
